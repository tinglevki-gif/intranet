import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def format_inline_text(paragraph, text, italic=False, color=None):
    parts = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        else:
            run = paragraph.add_run(part)
        
        if italic:
            run.italic = True
        if color:
            run.font.color.rgb = color

def create_tech_guide_docx(md_path, output_docx_path, base_img_dir):
    doc = Document()

    # Page Margins (2.0 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55) # Slate 700

    # Colors
    NAVY = RGBColor(0x0F, 0x17, 0x2A) # Slate 900
    SKY = RGBColor(0x02, 0x84, 0xC7)  # Sky 600
    DARK_BLUE = RGBColor(0x1E, 0x3A, 0x8A)
    CODE_BG = "0F172A" # Dark Slate 900

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_quote = False
    quote_text = []
    in_table = False
    table_lines = []

    for line in lines:
        line_str = line.rstrip('\r\n')
        stripped = line_str.strip()

        # Handle Code Blocks (``` ... ```)
        if stripped.startswith('```'):
            if in_code_block:
                # End of code block
                full_code = "\n".join(code_lines)
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                cell.width = Inches(6.8)
                set_cell_background(cell, CODE_BG)
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.05
                run = p.add_run(full_code)
                run.font.name = 'Consolas'
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0) # Slate 200
                
                doc.add_paragraph() # Spacing
                code_lines = []
                in_code_block = False
                continue
            else:
                in_code_block = True
                code_lines = []
                continue

        if in_code_block:
            code_lines.append(line_str)
            continue

        # Handle Tables (| ... |)
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            # End of table block -> render Word table
            if len(table_lines) >= 2:
                # Filter out separator row (e.g. |:---|:---|)
                raw_rows = [r for r in table_lines if not re.match(r'^\|[\s:-|]+\|$', r)]
                if raw_rows:
                    parsed_rows = [[c.strip() for c in r.strip('|').split('|')] for r in raw_rows]
                    cols_count = max(len(r) for r in parsed_rows)
                    word_table = doc.add_table(rows=len(parsed_rows), cols=cols_count)
                    word_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    
                    for row_idx, row_data in enumerate(parsed_rows):
                        for col_idx in range(cols_count):
                            cell_val = row_data[col_idx] if col_idx < len(row_data) else ""
                            cell = word_table.cell(row_idx, col_idx)
                            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.line_spacing = 1.05
                            
                            if row_idx == 0:
                                set_cell_background(cell, "0F172A") # Navy header
                                format_inline_text(p, cell_val, color=RGBColor(0xFF, 0xFF, 0xFF))
                                for r in p.runs:
                                    r.bold = True
                                    r.font.size = Pt(9)
                            else:
                                bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                                set_cell_background(cell, bg_color)
                                format_inline_text(p, cell_val, color=RGBColor(0x33, 0x41, 0x55))
                                for r in p.runs:
                                    r.font.size = Pt(8.5)
                    
                    doc.add_paragraph() # Spacing
            in_table = False
            table_lines = []

        # Handle Blockquotes
        if stripped.startswith('>'):
            in_quote = True
            quote_text.append(stripped.lstrip('> ').strip())
            continue
        elif in_quote:
            if quote_text:
                full_quote = " ".join(quote_text)
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                cell.width = Inches(6.8)
                set_cell_background(cell, "F0F9FF") # Light sky
                set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                format_inline_text(p, full_quote, italic=False, color=RGBColor(0x03, 0x69, 0xA1))
                doc.add_paragraph()
                quote_text = []
            in_quote = False

        if not stripped:
            continue

        # Horizontal Rule
        if stripped in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run('━' * 65)
            run.font.size = Pt(7)
            run.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
            continue

        # Headings
        if stripped.startswith('# '):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(stripped[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = DARK_BLUE
            continue
        elif stripped.startswith('## '):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(stripped[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = SKY
            continue
        elif stripped.startswith('### '):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            run = h.add_run(stripped[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = NAVY
            continue
        elif stripped.startswith('#### '):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            run = h.add_run(stripped[5:])
            run.font.name = 'Calibri'
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = SKY
            continue

        # Image tag: ![Alt text](image_path)
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            alt_text, rel_img_path = img_match.groups()
            img_filename = os.path.basename(rel_img_path)
            full_img_path = os.path.join(base_img_dir, img_filename)
            
            if os.path.exists(full_img_path):
                p_img = doc.add_paragraph()
                p_img.paragraph_format.space_before = Pt(6)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run_img = p_img.add_run()
                run_img.add_picture(full_img_path, width=Inches(6.2))
                
                p_cap = doc.add_paragraph()
                p_cap.paragraph_format.space_before = Pt(1)
                p_cap.paragraph_format.space_after = Pt(8)
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_run = p_cap.add_run(f"Abbildung: {alt_text}")
                cap_run.font.size = Pt(8.5)
                cap_run.font.italic = True
                cap_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            else:
                p_missing = doc.add_paragraph(f"[Diagramm: {alt_text} ({rel_img_path})]")
                p_missing.paragraph_format.space_after = Pt(6)
            continue

        # Bullet lists (- Item or * Item)
        if stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            format_inline_text(p, stripped[2:])
            continue

        # Numbered list (1. Item, 2. Item, etc.)
        num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if num_match:
            idx, content = num_match.groups()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            format_inline_text(p, content)
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        format_inline_text(p, stripped)

    os.makedirs(os.path.dirname(output_docx_path), exist_ok=True)
    doc.save(output_docx_path)
    print(f"Successfully generated: {output_docx_path}")

if __name__ == '__main__':
    md_file = 'docs/TECHNISCHES_BETRIEBSHANDBUCH.md'
    docx_file = 'docs/TECHNISCHES_BETRIEBSHANDBUCH.docx'
    img_dir = 'docs/images/tech_guide'
    create_tech_guide_docx(md_file, docx_file, img_dir)
